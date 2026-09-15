# ============================================================
# EXCEL REKAP TOOLS - VERSI 2.0
# ============================================================

import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# UTILITIES
# ============================================================

def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def extract_percentage(value):

    if value is None:
        return 0

    if isinstance(value, (int, float)):

        if 0 <= value <= 1:
            return value * 100

        return float(value)

    text = str(value).strip()

    match = re.search(
        r'(\d+(?:\.\d+)?)\s*%',
        text
    )

    if match:
        return float(match.group(1))

    try:

        number = float(text)

        if 0 <= number <= 1:
            return number * 100

        return number

    except:
        return 0


def read_components_from_row5(ws):

    components = {
        "PRS": 0,
        "Tugas1": 0,
        "Tugas2": 0,
        "Tugas3": 0,
        "QUIZ": 0,
        "Presnta": 0,
        "LAB": 0,
        "UTS": 0,
        "UAS": 0
    }

    for cell in ws[5]:

        value = cell.value

        if value is None:
            continue

        text = str(value).strip()

        percentage = extract_percentage(value)

        label = re.sub(
            r'\(\s*\d+(?:\.\d+)?\s*%\s*\)',
            '',
            text
        ).strip()

        normalized = re.sub(
            r'[^A-Z0-9]',
            '',
            label.upper()
        )

        if normalized == "PRS":
            components["PRS"] = percentage

        elif normalized in ["TUGAS1", "TGS1"]:
            components["Tugas1"] = percentage

        elif normalized in ["TUGAS2", "TGS2"]:
            components["Tugas2"] = percentage

        elif normalized in ["TUGAS3", "TGS3"]:
            components["Tugas3"] = percentage

        elif normalized in ["QUIZ", "KUIS"]:
            components["QUIZ"] = percentage

        elif normalized in [
            "PRESNTA",
            "PRESENTA",
            "PRESENTASI"
        ]:
            components["Presnta"] = percentage

        elif normalized == "LAB":
            components["LAB"] = percentage

        elif normalized == "UTS":
            components["UTS"] = percentage

        elif normalized == "UAS":
            components["UAS"] = percentage

    return components


# ============================================================
# FORMAT EXCEL
# ============================================================

def format_worksheet(ws):

    for cell in ws[1]:

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="D9EAF7"
        )

    thin = Side(
        style="thin",
        color="000000"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    for row in ws.iter_rows():

        for cell in row:

            cell.border = border

            cell.alignment = Alignment(
                vertical="center"
            )

    ws.freeze_panes = "A2"

    if ws.max_row > 1 and ws.max_column > 1:
        ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            try:

                length = len(
                    str(cell.value)
                )

                max_length = max(
                    max_length,
                    length
                )

            except:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            40
        )

    ws.row_dimensions[1].height = 30


# ============================================================
# MODUL 1
# REKAP MATA KULIAH & EVALUASI
# ============================================================

def process_excel(input_file, output_file):

    wb = load_workbook(
        input_file,
        data_only=False
    )

    output_wb = Workbook()

    default_sheet = output_wb.active

    output_wb.remove(default_sheet)

    # ========================================================
    # REKAP UTAMA
    # ========================================================

    ws_rekap = output_wb.create_sheet(
        "Rekap Utama"
    )

    headers = [
        "No",
        "Nama Kelas",
        "Kode Matkul",
        "Nama Matkul",
        "PRS (%)",
        "Tugas1 (%)",
        "Tugas2 (%)",
        "Tugas3 (%)",
        "QUIZ (%)",
        "Presnta (%)",
        "LAB (%)",
        "UTS (%)",
        "UAS (%)",
        "Total (%)",
        "Status"
    ]

    ws_rekap.append(headers)

    # ========================================================
    # EVALUASI DETAIL
    # ========================================================

    ws_eval = output_wb.create_sheet(
        "Evaluasi Detail"
    )

    eval_headers = [
        "No",
        "Nama Kelas",
        "Kode Matkul",
        "Nama Matkul",
        "Basis Evaluasi",
        "Komponen Evaluasi",
        "Bobot (%)",
        "Nama (Inggris)"
    ]

    ws_eval.append(eval_headers)

    nomor = 1
    nomor_eval = 1

    # ========================================================
    # BACA SEMUA SHEET
    # ========================================================

    for ws in wb.worksheets:

        nama_kelas = clean_text(
            ws["C6"].value
        )

        nama_matkul = clean_text(
            ws["D6"].value
        )

        kode_matkul = clean_text(
            ws["E6"].value
        )

        components = read_components_from_row5(ws)

        prs = components["PRS"]
        tugas1 = components["Tugas1"]
        tugas2 = components["Tugas2"]
        tugas3 = components["Tugas3"]
        quiz = components["QUIZ"]
        presnta = components["Presnta"]
        lab = components["LAB"]
        uts = components["UTS"]
        uas = components["UAS"]

        total = (
            prs +
            tugas1 +
            tugas2 +
            tugas3 +
            quiz +
            presnta +
            lab +
            uts +
            uas
        )

        status = (
            "VALID"
            if abs(total - 100) < 0.01
            else "CEK"
        )

        ws_rekap.append([
            nomor,
            nama_kelas,
            kode_matkul,
            nama_matkul,
            prs,
            tugas1,
            tugas2,
            tugas3,
            quiz,
            presnta,
            lab,
            uts,
            uas,
            total,
            status
        ])

        # ====================================================
        # KONVERSI EVALUASI
        # ====================================================

        tugas_total = (
            tugas1 +
            tugas2 +
            tugas3
        )

        project_total = (
            presnta +
            lab
        )

        evaluasi = [

            [
                2,
                "",
                prs,
                "Participatory Activity"
            ],

            [
                4,
                "TGS",
                tugas_total,
                "Assignment"
            ],

            [
                4,
                "QUIZ",
                quiz,
                "Quiz"
            ],

            [
                3,
                "",
                project_total,
                "Project Outcomes"
            ],

            [
                4,
                "UTS",
                uts,
                "Midterm Exam"
            ],

            [
                4,
                "UAS",
                uas,
                "Finalterm Exam"
            ]
        ]

        for item in evaluasi:

            ws_eval.append([
                nomor_eval,
                nama_kelas,
                kode_matkul,
                nama_matkul,
                item[0],
                item[1],
                item[2],
                item[3]
            ])

            nomor_eval += 1

        nomor += 1

    # ========================================================
    # FORMAT
    # ========================================================

    format_worksheet(ws_rekap)
    format_worksheet(ws_eval)

    # ========================================================
    # FORMAT ANGKA
    # ========================================================

    for row in ws_rekap.iter_rows(
        min_row=2,
        min_col=5,
        max_col=14
    ):

        for cell in row:
            cell.number_format = "0.00"

    for row in ws_eval.iter_rows(
        min_row=2,
        min_col=7,
        max_col=7
    ):

        for cell in row:
            cell.number_format = "0.00"

    # ========================================================
    # STATUS KONVERSI
    # ========================================================

    ws_eval["I1"] = "Status Konversi"

    row = 2

    while row <= ws_eval.max_row:

        nama_kelas = ws_eval.cell(
            row,
            2
        ).value

        kode_matkul = ws_eval.cell(
            row,
            3
        ).value

        total_eval = 0

        current = row

        while current <= ws_eval.max_row:

            if (
                ws_eval.cell(
                    current,
                    2
                ).value == nama_kelas
                and
                ws_eval.cell(
                    current,
                    3
                ).value == kode_matkul
            ):

                value = (
                    ws_eval.cell(
                        current,
                        7
                    ).value
                    or 0
                )

                total_eval += float(value)

                current += 1

            else:

                break

        status = (
            "VALID"
            if abs(total_eval - 100) < 0.01
            else "CEK"
        )

        ws_eval.cell(
            row,
            9
        ).value = status

        row = current

    output_wb.save(output_file)


# ============================================================
# MODUL 2
# REKAP MATA KULIAH & DOSEN
# ============================================================

def process_matkul_dosen(input_file, output_file):

    wb = load_workbook(
        input_file,
        data_only=True
    )

    output_wb = Workbook()

    ws = output_wb.active

    ws.title = "Rekap Mata Kuliah Dosen"

    headers = [
        "No",
        "Nama Sheet",
        "Kode Mata Kuliah",
        "Nama Mata Kuliah",
        "Nama Kelas",
        "Nama Dosen"
    ]

    ws.append(headers)

    nomor = 1

    for source_ws in wb.worksheets:

        try:

            nama_kelas = clean_text(
                source_ws["C6"].value
            )

            kode_mk = clean_text(
                source_ws["D6"].value
            )

            nama_mk = clean_text(
                source_ws["E6"].value
            )

            nama_dosen = clean_text(
                source_ws["L2"].value
            )

        except Exception:

            continue

        ws.append([
            nomor,
            source_ws.title,
            kode_mk,
            nama_mk,
            nama_kelas,
            nama_dosen
        ])

        nomor += 1

    format_worksheet(ws)

    output_wb.save(output_file)


# ============================================================
# MODUL 3
# MERGE SEMUA SHEET MULAI BARIS KE-5
# ============================================================

def process_merge_sheets(input_file, output_file):

    wb = load_workbook(
        input_file,
        data_only=True
    )

    output_wb = Workbook()

    ws_output = output_wb.active

    ws_output.title = "Merge Semua Sheet"

    first_sheet = True

    output_row = 1

    for source_ws in wb.worksheets:

        max_row = source_ws.max_row
        max_col = source_ws.max_column

        # Mulai dari baris ke-5
        for row_idx in range(5, max_row + 1):

            values = []

            for col_idx in range(1, max_col + 1):

                values.append(
                    source_ws.cell(
                        row_idx,
                        col_idx
                    ).value
                )

            # Tambahkan Nama Sheet
            values.append(source_ws.title)

            # Abaikan baris yang benar-benar kosong
            if all(
                value is None
                for value in values[:-1]
            ):
                continue

            for col_idx, value in enumerate(
                values,
                start=1
            ):

                ws_output.cell(
                    output_row,
                    col_idx
                ).value = value

            output_row += 1

    format_worksheet(ws_output)

    output_wb.save(output_file)


# ============================================================
# GUI APPLICATION
# ============================================================

class ExcelRekapApp:

    def __init__(self, root):

        self.root = root

        self.root.title(
            "Excel Rekap Tools"
        )

        self.root.geometry(
            "700x560"
        )

        self.root.resizable(
            False,
            False
        )

        self.input_file = ""

        self.selected_module = 1

        self.create_widgets()


    # ========================================================
    # GUI
    # ========================================================

    def create_widgets(self):

        title = tk.Label(
            self.root,
            text="EXCEL REKAP TOOLS",
            font=("Arial", 20, "bold")
        )

        title.pack(
            pady=(20, 5)
        )

        subtitle = tk.Label(
            self.root,
            text="Pengolahan & Rekap Data Excel",
            font=("Arial", 11)
        )

        subtitle.pack(
            pady=(0, 15)
        )

        module_frame = tk.LabelFrame(
            self.root,
            text="Pilih Modul",
            font=("Arial", 11, "bold"),
            padx=15,
            pady=10
        )

        module_frame.pack(
            padx=30,
            fill="x"
        )

        self.module_var = tk.IntVar(
            value=1
        )

        modules = [
            (
                1,
                "Rekap Mata Kuliah & Evaluasi"
            ),
            (
                2,
                "Rekap Mata Kuliah & Dosen"
            ),
            (
                3,
                "Merge Semua Sheet"
            )
        ]

        for value, text in modules:

            rb = tk.Radiobutton(
                module_frame,
                text=text,
                variable=self.module_var,
                value=value,
                command=self.module_changed,
                font=("Arial", 11),
                anchor="w"
            )

            rb.pack(
                fill="x",
                pady=4
            )

        frame = tk.Frame(
            self.root,
            padx=30
        )

        frame.pack(
            fill="both",
            expand=True
        )

        self.selected_label = tk.Label(
            frame,
            text="Modul aktif: Rekap Mata Kuliah & Evaluasi",
            font=("Arial", 12, "bold")
        )

        self.selected_label.pack(
            pady=(20, 12)
        )

        self.file_label = tk.Label(
            frame,
            text="Belum ada file Excel yang dipilih",
            anchor="w",
            relief="sunken",
            padx=10
        )

        self.file_label.pack(
            fill="x",
            ipady=8
        )

        btn_file = tk.Button(
            frame,
            text="PILIH FILE EXCEL",
            command=self.choose_file,
            width=25,
            height=2
        )

        btn_file.pack(
            pady=12
        )

        self.progress = ttk.Progressbar(
            frame,
            mode="indeterminate"
        )

        self.progress.pack(
            fill="x",
            pady=5
        )

        self.btn_process = tk.Button(
            frame,
            text="PROSES DATA",
            command=self.process,
            width=25,
            height=2,
            state="disabled"
        )

        self.btn_process.pack(
            pady=12
        )

        self.status_label = tk.Label(
            frame,
            text="Status: Siap",
            font=("Arial", 10)
        )

        self.status_label.pack(
            pady=5
        )

        footer = tk.Label(
            self.root,
            text="Excel Rekap Tools v2.0",
            font=("Arial", 9)
        )

        footer.pack(
            pady=12
        )


    # ========================================================
    # PILIH MODUL
    # ========================================================

    def module_changed(self):

        self.selected_module = self.module_var.get()

        module_names = {
            1: "Rekap Mata Kuliah & Evaluasi",
            2: "Rekap Mata Kuliah & Dosen",
            3: "Merge Semua Sheet"
        }

        self.selected_label.config(
            text="Modul aktif: "
            + module_names[
                self.selected_module
            ]
        )

        self.status_label.config(
            text="Status: Siap"
        )


    # ========================================================
    # PILIH FILE
    # ========================================================

    def choose_file(self):

        file = filedialog.askopenfilename(
            title="Pilih File Excel",
            filetypes=[
                (
                    "Excel Files",
                    "*.xlsx *.xlsm"
                ),
                (
                    "All Files",
                    "*.*"
                )
            ]
        )

        if file:

            self.input_file = file

            filename = os.path.basename(
                file
            )

            self.file_label.config(
                text=filename
            )

            self.btn_process.config(
                state="normal"
            )

            self.status_label.config(
                text="Status: File siap diproses"
            )


    # ========================================================
    # PROSES
    # ========================================================

    def process(self):

        if not self.input_file:

            messagebox.showwarning(
                "Peringatan",
                "Silakan pilih file Excel terlebih dahulu."
            )

            return

        initial_names = {
            1: "Rekap_Mata_Kuliah_dan_Evaluasi.xlsx",
            2: "Rekap_Mata_Kuliah_dan_Dosen.xlsx",
            3: "Merge_Semua_Sheet.xlsx"
        }

        output_file = filedialog.asksaveasfilename(
            title="Simpan Hasil Rekap",
            defaultextension=".xlsx",
            initialfile=initial_names[
                self.selected_module
            ],
            filetypes=[
                (
                    "Excel Files",
                    "*.xlsx"
                )
            ]
        )

        if not output_file:
            return

        try:

            self.status_label.config(
                text="Status: Sedang memproses..."
            )

            self.progress.start(10)

            self.btn_process.config(
                state="disabled"
            )

            self.root.update_idletasks()

            if self.selected_module == 1:

                process_excel(
                    self.input_file,
                    output_file
                )

            elif self.selected_module == 2:

                process_matkul_dosen(
                    self.input_file,
                    output_file
                )

            elif self.selected_module == 3:

                process_merge_sheets(
                    self.input_file,
                    output_file
                )

            self.progress.stop()

            self.btn_process.config(
                state="normal"
            )

            self.status_label.config(
                text="Status: Proses berhasil!"
            )

            messagebox.showinfo(
                "Berhasil",
                "Rekap Excel berhasil dibuat!\n\n"
                + output_file
            )

        except Exception as e:

            self.progress.stop()

            self.btn_process.config(
                state="normal"
            )

            self.status_label.config(
                text="Status: Terjadi kesalahan"
            )

            messagebox.showerror(
                "Error",
                "Terjadi kesalahan saat memproses file:\n\n"
                + str(e)
            )


# ============================================================
# START APPLICATION
# ============================================================

if __name__ == "__main__":

    root = tk.Tk()

    app = ExcelRekapApp(root)

    root.mainloop()
